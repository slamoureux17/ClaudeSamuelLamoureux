"""Importation d'un fichier « Scorecard Client One-Pager » (.xlsx) existant.

Le fichier est lu par libellés (et non par coordonnées fixes) pour tolérer
des lignes ajoutées ou déplacées.
"""
from __future__ import annotations

import io
import re

import openpyxl
from sqlalchemy.orm import Session

from ..models import ACTION_STATUSES, PROJETE, REEL, SERVICES, TEAM_ROLES, Action, Client
from ..repository import (
    get_or_create_client,
    set_services,
    set_team,
    upsert_assessment,
    upsert_metrics,
    upsert_review,
)
from ..textutil import normalize, parse_date, parse_number, parse_pct

SERVICE_BY_LABEL = {normalize(label): key for key, label in SERVICES}

TEAM_ROLE_BY_LABEL = {normalize(label): key for key, label in TEAM_ROLES}
STATUS_BY_LABEL = {normalize(label).split(" ")[0]: key for key, label in ACTION_STATUSES}

METRIC_ROWS = {
    "revenus totaux pieces et services": ("revenue", parse_number),
    "marge brute": ("gross_margin", parse_number),
    "% marge service": ("service_margin_pct", parse_pct),
    "% marge pieces": ("parts_margin_pct", parse_pct),
}


def _column_scenario(header: object) -> tuple[int, str] | None:
    """« 2024 » -> (2024, réel); « 2026 (projeté) » -> (2026, projeté); « 2026 Actuel » -> (2026, réel)."""
    match = re.search(r"(19|20)\d{2}", str(header or ""))
    if not match:
        return None
    year = int(match.group(0))
    return year, PROJETE if "proj" in normalize(header) else REEL


def _value(ws, ws_values, row: int, col: int):
    """Valeur saisie; pour une formule, la valeur calculée en cache par Excel."""
    value = ws.cell(row, col).value
    if isinstance(value, str) and value.startswith("="):
        return ws_values.cell(row, col).value
    return value


def import_onepager(session: Session, data: bytes) -> list[Client]:
    """Importe toutes les feuilles du classeur (une feuille = un client)."""
    wb = openpyxl.load_workbook(io.BytesIO(data))
    wb_values = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    clients = []
    for ws in wb.worksheets:
        client = _import_sheet(session, ws, wb_values[ws.title])
        if client:
            clients.append(client)
    session.flush()
    return clients


def _import_team(ws, client: Client) -> None:
    """Bloc « Équipe client » (colonne H: rôle, colonne I: nom), tel que produit par l'export."""
    header = next((r for r in range(1, ws.max_row + 1) if normalize(ws.cell(r, 8).value) == "equipe client"), None)
    if header is None:
        return
    main: dict[str, tuple[str | None, str | None]] = {}
    others: list[tuple[str | None, str | None]] = []
    row = header + 1
    while row <= ws.max_row and ws.cell(row, 8).value:
        label = str(ws.cell(row, 8).value).strip()
        name = str(ws.cell(row, 9).value or "").strip() or None
        title = re.search(r"\(([^)]+)\)\s*$", label)
        role = TEAM_ROLE_BY_LABEL.get(normalize(re.sub(r"\s*\([^)]*\)\s*$", "", label)))
        if role:
            main[role] = (name, title.group(1).strip() if title else None)
        else:
            others.append((label, name))
        row += 1
    set_team(client, main, others)


def _import_sheet(session: Session, ws, ws_values) -> Client | None:
    labels = {row: normalize(ws.cell(row, 1).value) for row in range(1, ws.max_row + 1)}
    find = lambda prefix: next((r for r, label in labels.items() if label.startswith(prefix)), None)  # noqa: E731

    name_row = find("nom du client")
    if not name_row or not ws.cell(name_row, 2).value:
        return None
    serti_row = find("num. adresse serti") or find("no. serti") or find("code serti")
    serti = str(ws.cell(serti_row, 2).value).strip() if serti_row and ws.cell(serti_row, 2).value else None
    client = get_or_create_client(session, str(ws.cell(name_row, 2).value).strip(), serti)
    branch_row = find("succursale")
    if branch_row and ws.cell(branch_row, 2).value:
        client.branch = str(ws.cell(branch_row, 2).value).strip()
    _import_team(ws, client)

    # --- Section 1: colonnes années / scénarios ---
    header_row = find("section") or find("1. donnees historiques")
    columns: dict[int, tuple[int, str]] = {}
    if header_row:
        for col in range(2, 6):
            parsed = _column_scenario(ws.cell(header_row, col).value)
            if parsed:
                columns[col] = parsed

    metric_values: dict[tuple[int, str], dict] = {}
    for row, label in labels.items():
        spec = METRIC_ROWS.get(label)
        if not spec:
            continue
        attr, parser = spec
        for col, key in columns.items():
            value = parser(_value(ws, ws_values, row, col))
            if value is not None:
                metric_values.setdefault(key, {})[attr] = value
    for (year, scenario), values in metric_values.items():
        upsert_metrics(session, client, year, scenario, "excel", **values)

    services_start = find("nombre de services")
    if services_start:
        used: dict[tuple[int, str], set[str]] = {key: set() for key in columns.values()}
        row = services_start + 1
        while row <= ws.max_row and labels.get(row) in SERVICE_BY_LABEL:
            for col, key in columns.items():
                if str(ws.cell(row, col).value or "").strip():
                    used[key].add(SERVICE_BY_LABEL[labels[row]])
            row += 1
        for (year, scenario), services in used.items():
            set_services(session, client, year, scenario, services)

    # --- Section 2: potentiel ---
    potential_row = find("potentiel annuel")
    if potential_row:
        client.potential_annual = parse_number(_value(ws, ws_values, potential_row, 2))
    current_row = find("revenus actuels")
    if current_row:
        client.current_revenue_override = parse_number(_value(ws, ws_values, current_row, 2))

    # --- Section 3: analyse qualitative ---
    qual_row = find("3. analyse qualitative")
    if qual_row:
        qual_years = {
            col: int(ws.cell(qual_row, col).value)
            for col in range(2, 6)
            if str(ws.cell(qual_row, col).value or "").strip().isdigit()
        }
        top_row, sat_row = find("resultat topgrading"), find("niveau de satisfaction")
        for col, year in qual_years.items():
            values = {}
            if top_row and ws.cell(top_row, col).value:
                values["topgrading"] = str(ws.cell(top_row, col).value).strip().upper()[:2]
            if sat_row and parse_number(ws.cell(sat_row, col).value) is not None:
                values["satisfaction"] = int(parse_number(ws.cell(sat_row, col).value))
            if values:
                upsert_assessment(session, client, year, **values)
        # Les champs texte (cellules fusionnées) décrivent la situation actuelle -> année la plus récente.
        if qual_years:
            text = {}
            for prefix, attr in (("points forts", "strengths"), ("enjeux", "weaknesses"),
                                 ("potentiel de croissance", "growth_potential")):
                row = find(prefix)
                if row and ws.cell(row, 2).value:
                    text[attr] = str(ws.cell(row, 2).value).strip()
            if text:
                upsert_assessment(session, client, max(qual_years.values()), **text)

    # --- Section 4: plan d'action ---
    existing_titles = {normalize(a.title) for a in client.actions}
    for row, label in labels.items():
        match = re.fullmatch(r"action (\d+)", label)
        if not match or not ws.cell(row, 2).value:
            continue
        n = match.group(1)
        title = str(ws.cell(row, 2).value).strip()
        if normalize(title) in existing_titles:
            continue
        field = lambda prefix: next(  # noqa: E731
            (ws.cell(r, 2).value for r in range(row + 1, min(row + 5, ws.max_row + 1)) if labels[r] == f"{prefix} {n}"),
            None,
        )
        follow_up = str(field("suivi") or "").strip() or None
        status = None
        # Format de l'export: « [En cours] texte du suivi »
        tagged = re.match(r"\[([^\]]+)\]\s*", follow_up or "")
        if tagged:
            status = STATUS_BY_LABEL.get(normalize(tagged.group(1)).split(" ")[0])
            follow_up = follow_up[tagged.end():] or None
        done = status == "fait" or bool(follow_up and normalize(follow_up).startswith("fait"))
        client.actions.append(Action(
            title=title,
            owner=str(field("responsable") or "").strip() or None,
            due_date=parse_date(field("echeance")),
            follow_up=follow_up,
            status="fait" if done else (status or "a_faire"),
            done_date=parse_date(tagged.group(1) if tagged else follow_up) if done else None,
        ))

    # --- Section 5: suivi trimestriel (colonnes H/I/J) ---
    current_year = max((y for y, _ in columns.values()), default=None)
    if current_year:
        for row in range(1, ws.max_row + 1):
            match = re.fullmatch(r"t([1-4])", normalize(ws.cell(row, 8).value))
            if not match:
                continue
            result, notes = ws.cell(row, 9).value, ws.cell(row, 10).value
            if result or notes:
                upsert_review(session, client, current_year, int(match.group(1)),
                              result=str(result).strip() if result else None,
                              notes=str(notes).strip() if notes else None)
    return client
