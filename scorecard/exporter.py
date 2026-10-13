"""Export d'un client au format du gabarit « Scorecard Client One-Pager » (.xlsx)."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .metrics import build_scorecard
from .models import ACTION_STATUS_LABELS, Client

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
WHITE_BOLD = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
THIN = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
WRAP = Alignment(wrap_text=True, vertical="top")
MONEY = '#,##0 $;-#,##0 $'
PCT = "0.0%"


def export_client(client: Client, year: int) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Scorecard Client"
    write_sheet(ws, client, year)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def write_sheet(ws, client: Client, year: int) -> None:
    sc = build_scorecard(client, year)
    cols = sc["columns"]

    ws["A1"], ws["B1"] = "Nom du client", client.name
    ws["A2"], ws["B2"] = "Num. adresse Serti", client.serti_code
    ws["A1"].font = ws["A2"].font = BOLD

    # En-tête section 1 (B..E = colonnes, F = différence)
    header = ["Section / Élément"] + [c.label for c in cols] + ["Différence"]
    for i, value in enumerate(header, start=1):
        cell = ws.cell(4, i, value)
        cell.font, cell.fill = WHITE_BOLD, HEADER_FILL
    _section(ws, 5, "1. Données historiques", 6)

    row = 6
    metric_rows = {}
    for r in sc["rows"]:
        ws.cell(row, 1, r["label"])
        metric_rows[r["field"]] = row
        for i, col in enumerate(cols, start=2):
            letter = get_column_letter(i)
            if r["field"] == "global_pct":
                # Formule vivante comme dans le gabarit
                cell = ws.cell(row, i, f"=IF({letter}{metric_rows['revenue']}=0,\"\","
                                       f"{letter}{metric_rows['gross_margin']}/{letter}{metric_rows['revenue']})")
            else:
                cell = ws.cell(row, i, r["values"][col.key])
            cell.number_format = MONEY if r["fmt"] == "money" else PCT
        if r["field"] in ("revenue", "gross_margin"):
            ws.cell(row, 6, f"=IF(OR(D{row}=\"\",E{row}=\"\"),\"\",D{row}-E{row})").number_format = MONEY
        row += 1

    ws.cell(row, 1, "Nombre de services utilisés").font = BOLD
    for i, col in enumerate(cols, start=2):
        ws.cell(row, i, f"=COUNTIF({get_column_letter(i)}{row + 1}:{get_column_letter(i)}{row + len(sc['service_list'])},\"X\")")
    row += 1
    for key, label in sc["service_list"]:
        ws.cell(row, 1, f"   {label}")
        for i, col in enumerate(cols, start=2):
            if key in sc["services"][col.key]:
                ws.cell(row, i, "X").alignment = Alignment(horizontal="center")
        row += 1

    # Section 2
    row += 1
    _section(ws, row, "2. Potentiel global du client", 6)
    pot_row = row + 1
    ws.cell(pot_row, 1, "Potentiel annuel total (estimé)")
    ws.cell(pot_row, 2, sc["potential"].potential).number_format = MONEY
    ws.cell(pot_row + 1, 1, "Revenus actuels avec Camions BL")
    ws.cell(pot_row + 1, 2, sc["potential"].current).number_format = MONEY
    ws.cell(pot_row + 2, 1, "% de pénétration actuelle")
    ws.cell(pot_row + 2, 2, f"=IF(OR(B{pot_row}=\"\",B{pot_row}=0),\"\",B{pot_row + 1}/B{pot_row})").number_format = PCT
    ws.cell(pot_row + 3, 1, "Gap à combler")
    ws.cell(pot_row + 3, 2, f"=IF(B{pot_row}=\"\",\"\",B{pot_row}-B{pot_row + 1})").number_format = MONEY
    row = pot_row + 5

    # Section 3
    _section(ws, row, "3. Analyse qualitative", 6)
    for i, y in enumerate(sc["assessment_years"], start=2):
        ws.cell(row, i, y).font = BOLD
    assessments = sc["assessments"]
    ws.cell(row + 1, 1, "Résultat Topgrading Client")
    ws.cell(row + 2, 1, "Niveau de satisfaction client (/5)")
    for i, y in enumerate(sc["assessment_years"], start=2):
        a = assessments[y]
        if a:
            ws.cell(row + 1, i, a.topgrading)
            ws.cell(row + 2, i, a.satisfaction)
    latest = next((assessments[y] for y in reversed(sc["assessment_years"])
                   if assessments[y] and (assessments[y].strengths or assessments[y].weaknesses
                                          or assessments[y].growth_potential)), None)
    row += 3
    for label, attr in (("Points forts de la relation", "strengths"), ("Enjeux ou points faibles", "weaknesses"),
                        ("Potentiel de croissance", "growth_potential")):
        ws.cell(row, 1, label)
        _merged_text(ws, row, getattr(latest, attr) if latest else None)
        row += 1

    # Section 4
    row += 2
    _section(ws, row, "4. Plan d’action", 6)
    row += 1
    for n, action in enumerate(client.actions, start=1):
        for label, value in ((f"Action {n}", action.title), (f"Responsable {n}", action.owner),
                             (f"Échéance {n}", action.due_date.strftime("%m/%d/%Y") if action.due_date else None),
                             (f"Suivi {n}", _follow_up(action))):
            ws.cell(row, 1, label).font = BOLD if label.startswith("Action") else Font()
            _merged_text(ws, row, value)
            row += 1
        row += 1

    # Section 5 (colonnes H..J)
    for i, value in enumerate(["5. Suivi de performance", "Résultats", "Notes"], start=8):
        cell = ws.cell(4, i, value)
        cell.font, cell.fill = WHITE_BOLD, HEADER_FILL
    for q in range(1, 5):
        review = sc["reviews"][q]
        ws.cell(4 + q, 8, f"T{q}").font = BOLD
        ws.cell(4 + q, 9, review.result if review else None).alignment = WRAP
        ws.cell(4 + q, 10, review.notes if review else None).alignment = WRAP

    for letter, width in {"A": 38, "B": 16, "C": 16, "D": 16, "E": 16, "F": 14, "G": 3, "H": 24, "I": 40, "J": 60}.items():
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "B5"


def _section(ws, row: int, title: str, width: int) -> None:
    for col in range(1, width + 1):
        ws.cell(row, col).fill = SECTION_FILL
    ws.cell(row, 1, title).font = BOLD


def _merged_text(ws, row: int, value) -> None:
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
    cell = ws.cell(row, 2, value)
    cell.alignment = WRAP
    if value and len(str(value)) > 70:
        ws.row_dimensions[row].height = 15 * (len(str(value)) // 70 + 1)


def _follow_up(action) -> str | None:
    status = ACTION_STATUS_LABELS[action.status]
    if action.status == "fait" and action.done_date:
        status += f" {action.done_date.strftime('%m/%d/%Y')}"
    if action.follow_up and action.follow_up.lower().startswith(status.lower()[:4]):
        return action.follow_up
    return f"[{status}] {action.follow_up}" if action.follow_up else f"[{status}]"
