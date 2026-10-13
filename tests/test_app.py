from pathlib import Path

import openpyxl
import io
import pytest
from fastapi.testclient import TestClient

from scorecard import db
from scorecard.app import app

EXAMPLE = Path(__file__).parents[1] / "exemples"


@pytest.fixture
def web(tmp_path, monkeypatch):
    monkeypatch.setenv("SCORECARD_YEAR", "2026")
    monkeypatch.setenv("PIPEDRIVE_API_TOKEN", "")
    db.init_db(f"sqlite:///{tmp_path / 'test.db'}")
    with TestClient(app) as client:
        yield client


def test_import_edit_and_export(web):
    with open(EXAMPLE / "Scorecard_Client_OnePager.xlsx", "rb") as f:
        r = web.post("/importer/onepager", files={"fichier": ("s.xlsx", f)}, follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"].startswith("/clients/1")

    page = web.get("/clients/1?year=2026")
    assert "JMF Transport" in page.text and "850" in page.text

    form = {"year": "2026", "name": "JMF Transport", "serti_code": "TRAJMF",
            "m__2026_reel__revenue": "60 000", "m__2026_reel__gross_margin": "15000",
            "m__2026_projete__revenue": "180000", "svc__2026_reel__jockey": "1",
            "potential_annual": "900000", "a__2026__topgrading": "B", "a__2026__satisfaction": "5",
            "strengths": "Bonne relation", "r__2__result": "Objectif T2 atteint",
            "branch": "Drummondville", "team__experience_client__name": "Julie Côté",
            "team__ventes_camions__name": "Marc Roy", "team__direction__name": "Samuel L.",
            "team__direction__title": "Co-Président", "team__conseiller_technique__name": "Éric B.",
            "other_count": "2", "other__0__title": "Représentant pièces", "other__0__name": "Pascal",
            "other__1__title": "", "other__1__name": ""}
    r = web.post("/clients/1/fiche", data=form, follow_redirects=False)
    assert r.status_code == 303

    r = web.post("/clients/1/actions", data={"title": "Proposer Excelerator", "owner": "Maxime",
                                             "due_date": "2026-12-15", "year": "2026"})
    assert r.status_code == 200

    dash = web.get("/")
    assert "JMF Transport" in dash.text and "Drummondville" in dash.text and "Julie Côté" in dash.text
    assert "JMF Transport" in web.get("/?succursale=Drummondville").text
    assert "JMF Transport" not in web.get("/?succursale=Laval").text
    assert "JMF Transport" in web.get("/?q=julie").text

    page = web.get("/clients/1")
    assert 'value="Marc Roy"' in page.text and "<option selected>Co-Président</option>" in page.text

    xlsx = web.get("/clients/1/export.xlsx?year=2026")
    ws = openpyxl.load_workbook(io.BytesIO(xlsx.content)).active
    assert ws["B1"].value == "JMF Transport"
    assert ws["E6"].value == 60000 and ws["D6"].value == 180000
    values = [c.value for row in ws.iter_rows() for c in row if c.value]
    assert "Proposer Excelerator" in values and "Objectif T2 atteint" in values and "B" in values

    # L'export se réimporte (aller-retour)
    from scorecard.importers.onepager import import_onepager
    from sqlalchemy import create_engine
    from sqlalchemy.orm import Session
    from scorecard.models import Base
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    with Session(engine) as s:
        [client] = import_onepager(s, xlsx.content)
        assert client.potential_annual == 900000
        assert client.branch == "Drummondville" and client.cx_rep == "Julie Côté"
        direction = client.team_member("direction")
        assert (direction.name, direction.title) == ("Samuel L.", "Co-Président")
        assert [(m.title, m.name) for m in client.other_team] == [("Représentant pièces", "Pascal")]
        assert len(client.actions) == 4
        assert client.actions[0].follow_up.startswith("à planifier")
        assert client.actions[1].status == "fait"


def test_clearing_a_name_removes_the_member(web):
    web.post("/clients", data={"name": "Test", "branch": "Laval", "cx_rep": "Julie"})
    page = web.get("/clients/1")
    assert 'value="Julie"' in page.text
    web.post("/clients/1/fiche", data={"year": "2026", "name": "Test", "team__experience_client__name": ""})
    assert "Julie" not in web.get("/").text


def test_migrates_old_database(tmp_path):
    import sqlite3

    path = tmp_path / "old.db"
    con = sqlite3.connect(path)
    con.execute("CREATE TABLE clients (id INTEGER PRIMARY KEY, name VARCHAR(200), serti_code VARCHAR(50), "
                "account_manager VARCHAR(100), potential_annual FLOAT, current_revenue_override FLOAT, "
                "pipedrive_org_id INTEGER, pipedrive_snapshot JSON, pipedrive_synced_at DATETIME, "
                "created_at DATETIME, updated_at DATETIME)")
    con.execute("INSERT INTO clients (id, name, account_manager) VALUES (1, 'Vieux client', 'Maxime')")
    con.commit()
    con.close()
    db.init_db(f"sqlite:///{path}")
    with db._SessionLocal() as s:
        from scorecard.models import Client
        client = s.get(Client, 1)
        assert client.branch is None
        assert [(m.title, m.name) for m in client.other_team] == [("Responsable du compte", "Maxime")]
    db.init_db(f"sqlite:///{path}")  # idempotent
    with db._SessionLocal() as s:
        assert len(s.get(Client, 1).team) == 1


def test_pipedrive_disabled_shows_hint(web):
    web.post("/clients", data={"name": "Test"})
    page = web.get("/clients/1")
    assert "PIPEDRIVE_API_TOKEN" in page.text
